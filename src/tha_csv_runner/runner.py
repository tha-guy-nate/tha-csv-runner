import csv
import functools
import json
import shutil
from collections.abc import Callable
from datetime import datetime
from pathlib import Path
from typing import Any

import openpyxl
from tqdm import tqdm

from .errors import CsvError


def tqdm_ncols(max_cols: int = 85) -> int:
    return min(shutil.get_terminal_size(fallback=(max_cols, 24)).columns, max_cols)


def _sort_key(val: object) -> tuple[int, float | str]:
    try:
        return (0, float(val))  # type: ignore[arg-type]
    except (TypeError, ValueError):
        return (1, str(val))


def _is_excel_path(path: Path) -> bool:
    return path.suffix.lower() == ".xlsx"


def _is_jsonl_path(path: Path) -> bool:
    return path.suffix.lower() == ".jsonl"


def _read_jsonl_rows(path: Path, encoding: str) -> tuple[list[str] | None, list[dict[str, Any]]]:
    raw_rows = []
    with open(path, encoding=encoding) as f:
        for line in f:
            line = line.strip()
            if line:
                raw_rows.append(json.loads(line))
    fieldnames = list(raw_rows[0].keys()) if raw_rows else None
    return fieldnames, raw_rows


def _resolve_sheet(wb: Any, path: Path, sheet: str | int | None) -> Any:
    if sheet is None:
        return wb.active
    if isinstance(sheet, int):
        try:
            return wb.worksheets[sheet]
        except IndexError:
            raise CsvError(
                f"Sheet index {sheet} out of range for {path} ({len(wb.worksheets)} sheet(s))"
            ) from None
    try:
        return wb[sheet]
    except KeyError:
        raise CsvError(
            f"Sheet '{sheet}' not found in {path}. Available sheets: {wb.sheetnames}"
        ) from None


def _read_excel_rows(
    path: Path, sheet: str | int | None = None
) -> tuple[list[str] | None, list[dict[str, Any]]]:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        ws = _resolve_sheet(wb, path, sheet)
        rows_iter = ws.iter_rows(values_only=True)
        try:
            header_row = next(rows_iter)
        except StopIteration:
            return None, []
        fieldnames = [str(h) if h is not None else "" for h in header_row]
        raw_rows = [dict(zip(fieldnames, row, strict=False)) for row in rows_iter]
    finally:
        wb.close()
    return fieldnames, raw_rows


def _write_chunk_xlsx(
    path: Path,
    rows: list[dict[str, Any]],
    cols: list[str],
    label: str,
    sheet: str | None = None,
    show_progress: bool = True,
) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = openpyxl.Workbook()
    ws = wb.active
    if sheet is not None:
        ws.title = sheet
    if rows:
        ws.append(cols)
        for row in tqdm(rows, desc=label, ncols=tqdm_ncols(), disable=not show_progress):
            ws.append([row.get(c) for c in cols])
    wb.save(path)


def _write_chunk_jsonl(
    path: Path,
    rows: list[dict[str, Any]],
    cols: list[str],
    label: str,
    encoding: str = "utf-8",
    show_progress: bool = True,
) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with open(path, "w", encoding=encoding) as f:
        for row in tqdm(rows, desc=label, ncols=tqdm_ncols(), disable=not show_progress):
            f.write(json.dumps({c: row[c] for c in cols if c in row}) + "\n")


def _write_chunk(
    path: Path,
    rows: list[dict[str, Any]],
    cols: list[str],
    label: str,
    delimiter: str = ",",
    encoding: str = "utf-8",
    sheet: str | None = None,
    show_progress: bool = True,
) -> None:
    if _is_excel_path(path):
        _write_chunk_xlsx(path, rows, cols, label, sheet, show_progress)
        return
    if _is_jsonl_path(path):
        _write_chunk_jsonl(path, rows, cols, label, encoding, show_progress)
        return
    path.parent.mkdir(parents=True, exist_ok=True)
    with open(path, "w", newline="", encoding=encoding) as f:
        if rows:
            writer = csv.DictWriter(f, fieldnames=cols, delimiter=delimiter)
            writer.writeheader()
            writer.writerows(
                {c: row[c] for c in cols if c in row}
                for row in tqdm(rows, desc=label, ncols=tqdm_ncols(), disable=not show_progress)
            )


class ThaCSV:
    def __init__(
        self, delimiter: str = ",", encoding: str = "utf-8", show_progress: bool = True
    ) -> None:
        self.rows: list[dict[str, Any]] = []
        self._read: bool = False
        self._input_path: Path | None = None
        self._delimiter = delimiter
        self._encoding = encoding
        self._show_progress = show_progress
        self.status_cb = print

    def read(
        self,
        desc: str | None,
        input_path: str | Path,
        required_headers: list[str],
        validator: Callable[[dict[str, Any]], None] | None = None,
        enrich: bool = True,
        sheet: str | int | None = None,
    ) -> list[dict[str, Any]]:
        self._input_path = Path(input_path)

        if sheet is not None and not _is_excel_path(self._input_path):
            raise ValueError("sheet= is only valid when reading a .xlsx file")

        if _is_excel_path(self._input_path):
            fieldnames, raw_rows = _read_excel_rows(self._input_path, sheet)
        elif _is_jsonl_path(self._input_path):
            fieldnames, raw_rows = _read_jsonl_rows(self._input_path, self._encoding)
        else:
            with open(self._input_path, newline="", encoding=self._encoding) as f:
                reader = csv.DictReader(f, delimiter=self._delimiter)
                fieldnames = list(reader.fieldnames) if reader.fieldnames is not None else None
                raw_rows = list(reader)

        if fieldnames is None:
            raise CsvError(f"{self._input_path} appears to be empty")
        missing = [h for h in required_headers if h not in fieldnames]
        if missing:
            raise CsvError(f"Missing required headers: {missing}")

        self.rows = []
        self._read = True

        reading = f"Reading {self._input_path.stem} CSV"
        label = f"{desc}: {reading}" if desc is not None else reading
        for i, row in enumerate(
            tqdm(raw_rows, desc=label, ncols=tqdm_ncols(), disable=not self._show_progress),
            start=2,
        ):
            if enrich:
                enriched = {**row, "row number": i, "row status": "", "message": ""}
            else:
                enriched = dict(row)
            try:
                if validator is not None:
                    validator(enriched)
            except Exception as exc:
                if enrich:
                    enriched["row status"] = "error"
                    enriched["message"] = str(exc)
                else:
                    raise
            self.rows.append(enriched)

        return self.rows

    def write(
        self,
        desc: str | None,
        output_path: str | Path | None = None,
        rows: list[dict[str, Any]] | None = None,
        sort_by: str | list[str] | None = None,
        ascending: bool | list[bool] = True,
        column_order: list[str] | None = None,
        keep: list[str] | None = None,
        drop: list[str] | None = None,
        chunk_size: int | None = None,
        sheet: str | None = None,
    ) -> Path | list[Path]:
        if rows is None and not self._read:
            raise RuntimeError("No data to write — call read() first or pass rows=")
        if keep and drop:
            raise ValueError("Cannot specify both keep and drop")
        if chunk_size is not None and chunk_size < 1:
            raise ValueError("chunk_size must be >= 1")

        rows = list(rows) if rows is not None else list(self.rows)

        # --- column filtering ---
        # Union of all rows' keys, not just rows[0] — a "narrow" first row (e.g. from an
        # error/no-match path) must not silently drop columns that later rows carry.
        all_cols: list[str] = []
        seen: set[str] = set()
        for row in rows:
            for key in row:
                if key not in seen:
                    seen.add(key)
                    all_cols.append(key)

        if keep:
            cols = [c for c in keep if c in all_cols]
        elif drop:
            cols = [c for c in all_cols if c not in drop]
        else:
            cols = all_cols

        # --- column ordering: listed cols first, unlisted follow in original order ---
        if column_order:
            front = [c for c in column_order if c in cols]
            rest = [c for c in cols if c not in column_order]
            cols = front + rest

        # --- sorting ---
        if sort_by is not None:
            sort_cols = [sort_by] if isinstance(sort_by, str) else list(sort_by)
            asc_list = (
                [ascending] * len(sort_cols) if isinstance(ascending, bool) else list(ascending)
            )

            def compare(a: dict[str, Any], b: dict[str, Any]) -> int:
                for col, asc in zip(sort_cols, asc_list, strict=True):
                    ka, kb = _sort_key(a.get(col, "")), _sort_key(b.get(col, ""))
                    if ka < kb:
                        return -1 if asc else 1
                    if ka > kb:
                        return 1 if asc else -1
                return 0

            rows.sort(key=functools.cmp_to_key(compare))

        # --- output path ---
        if output_path is None:
            stem = self._input_path.stem if self._input_path else "output"
            ts = datetime.now().strftime("%Y%m%d_%H%M%S")
            output_path = Path(f"{stem}_processed_{ts}.csv")

        output_file = Path(output_path)

        if sheet is not None and not _is_excel_path(output_file):
            raise ValueError("sheet= is only valid when writing a .xlsx file")

        # --- chunked write ---
        if chunk_size is not None:
            chunks = [rows[i : i + chunk_size] for i in range(0, max(len(rows), 1), chunk_size)]
            paths = []
            for idx, chunk in enumerate(chunks, start=1):
                chunk_name = f"{output_file.stem}_{idx:03d}{output_file.suffix}"
                chunk_path = output_file.parent / chunk_name
                writing = f"Writing {output_file.stem} CSV ({idx}/{len(chunks)})"
                label = f"{desc} ({idx}/{len(chunks)}): {writing}" if desc else writing
                _write_chunk(
                    chunk_path,
                    chunk,
                    cols,
                    label,
                    self._delimiter,
                    self._encoding,
                    sheet,
                    self._show_progress,
                )
                paths.append(chunk_path)
            self.status_cb(f"✅ Done! CSV was written to: {paths}")
            return paths

        writing = f"Writing {output_file.stem} CSV"
        write_label = f"{desc}: {writing}" if desc is not None else writing
        _write_chunk(
            output_file,
            rows,
            cols,
            write_label,
            self._delimiter,
            self._encoding,
            sheet,
            self._show_progress,
        )
        self.status_cb(f"✅ Done! CSV was written to: {output_file}")
        return output_file
