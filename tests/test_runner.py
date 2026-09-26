import csv
from pathlib import Path

import pytest

from tha_csv_runner import ThaCSV
from tha_csv_runner.errors import CsvError


def noop(row: dict) -> None:
    pass


def fail_on_bob(row: dict) -> None:
    if row["name"] == "Bob":
        raise ValueError("Bob is not allowed")


def test_empty_file_raises(tmp_path: Path) -> None:
    csv_path = tmp_path / "empty.csv"
    csv_path.write_text("")
    runner = ThaCSV()
    with pytest.raises(CsvError, match="appears to be empty"):
        runner.read(None, csv_path, ["name"])


def test_happy_path(simple_csv: Path) -> None:
    runner = ThaCSV()
    runner.read(None, simple_csv, ["id", "name", "email"], noop)
    assert len(runner.rows) == 3
    assert all(r["row status"] == "" for r in runner.rows)


def test_read_returns_rows(simple_csv: Path) -> None:
    runner = ThaCSV()
    result = runner.read(None, simple_csv, ["name"])
    assert result is runner.rows
    assert len(result) == 3


def test_row_number_injected(simple_csv: Path) -> None:
    runner = ThaCSV()
    runner.read(None, simple_csv, ["name"])
    assert runner.rows[0]["row number"] == 2
    assert runner.rows[2]["row number"] == 4


def test_message_and_status_columns_present(simple_csv: Path) -> None:
    runner = ThaCSV()
    runner.read(None, simple_csv, ["name"])
    for row in runner.rows:
        assert "row number" in row
        assert "row status" in row
        assert "message" in row


def test_error_row_captured(simple_csv: Path) -> None:
    runner = ThaCSV()
    runner.read(None, simple_csv, ["name"], fail_on_bob)

    errors = [r for r in runner.rows if r["row status"] == "error"]
    success = [r for r in runner.rows if r["row status"] == ""]

    assert len(errors) == 1
    assert errors[0]["name"] == "Bob"
    assert "Bob is not allowed" in errors[0]["message"]
    assert len(success) == 2


def test_missing_required_header_raises(simple_csv: Path) -> None:
    runner = ThaCSV()
    with pytest.raises(CsvError, match="Missing required headers"):
        runner.read(None, simple_csv, ["id", "phone"])


def test_original_columns_preserved(simple_csv: Path) -> None:
    runner = ThaCSV()
    runner.read(None, simple_csv, ["name", "email"])
    assert runner.rows[0]["name"] == "Alice"
    assert runner.rows[0]["email"] == "alice@example.com"


def test_write_creates_file(simple_csv: Path, tmp_path: Path) -> None:
    out = tmp_path / "out.csv"
    runner = ThaCSV()
    runner.read(None, simple_csv, ["name"])
    runner.write(None, out)
    assert out.exists()


def test_write_contains_all_rows(simple_csv: Path, tmp_path: Path) -> None:
    out = tmp_path / "out.csv"
    runner = ThaCSV()
    runner.read(None, simple_csv, ["name"])
    runner.write(None, out)
    rows = list(csv.DictReader(out.open()))
    assert len(rows) == 3


def test_write_includes_enriched_columns(simple_csv: Path, tmp_path: Path) -> None:
    out = tmp_path / "out.csv"
    runner = ThaCSV()
    runner.read(None, simple_csv, ["name"])
    runner.write(None, out)
    rows = list(csv.DictReader(out.open()))
    assert "row number" in rows[0]
    assert "row status" in rows[0]
    assert "message" in rows[0]


def test_write_auto_names_file(
    simple_csv: Path, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.chdir(tmp_path)
    runner = ThaCSV()
    runner.read(None, simple_csv, ["name"])
    out = runner.write(None)
    assert out.name.startswith("simple_processed_")
    assert out.exists()


def test_write_before_read_raises(simple_csv: Path) -> None:
    runner = ThaCSV()
    with pytest.raises(RuntimeError, match=r"call read\(\)"):
        runner.write(None)


def test_write_rows_param_bypasses_read_guard(tmp_path: Path) -> None:
    out = tmp_path / "out.csv"
    rows = [{"name": "Alice", "val": "1"}, {"name": "Bob", "val": "2"}]
    runner = ThaCSV()
    result = runner.write(None, out, rows=rows)
    assert isinstance(result, Path)
    written = list(csv.DictReader(out.open()))
    assert len(written) == 2
    assert written[0]["name"] == "Alice"


def test_write_sort_by_single(simple_csv: Path, tmp_path: Path) -> None:
    out = tmp_path / "out.csv"
    runner = ThaCSV()
    runner.read(None, simple_csv, ["name"])
    runner.write(None, out, sort_by="name", ascending=False)
    rows = list(csv.DictReader(out.open()))
    names = [r["name"] for r in rows]
    assert names == sorted(names, reverse=True)


def test_write_sort_by_multiple(simple_csv: Path, tmp_path: Path) -> None:
    out = tmp_path / "out.csv"
    runner = ThaCSV()
    runner.read(None, simple_csv, ["name"])
    runner.write(None, out, sort_by=["name", "email"], ascending=[True, False])
    rows = list(csv.DictReader(out.open()))
    assert len(rows) == 3


def test_write_keep(simple_csv: Path, tmp_path: Path) -> None:
    out = tmp_path / "out.csv"
    runner = ThaCSV()
    runner.read(None, simple_csv, ["name"])
    runner.write(None, out, keep=["name", "email"])
    rows = list(csv.DictReader(out.open()))
    assert list(rows[0].keys()) == ["name", "email"]


def test_write_drop(simple_csv: Path, tmp_path: Path) -> None:
    out = tmp_path / "out.csv"
    runner = ThaCSV()
    runner.read(None, simple_csv, ["name"])
    runner.write(None, out, drop=["row number", "row status", "message"])
    rows = list(csv.DictReader(out.open()))
    assert "row number" not in rows[0]
    assert "name" in rows[0]


def test_write_header_is_union_of_all_rows(tmp_path: Path) -> None:
    out = tmp_path / "out.csv"
    runner = ThaCSV()
    runner.write(None, out, rows=[{"a": 1}, {"a": 2, "b": 3}])
    rows = list(csv.DictReader(out.open()))
    assert list(rows[0].keys()) == ["a", "b"]
    assert rows[0]["b"] == ""
    assert rows[1]["b"] == "3"


def test_write_keep_uses_union_of_all_rows(tmp_path: Path) -> None:
    out = tmp_path / "out.csv"
    runner = ThaCSV()
    runner.write(None, out, rows=[{"a": 1}, {"a": 2, "b": 3}], keep=["a", "b"])
    rows = list(csv.DictReader(out.open()))
    assert list(rows[0].keys()) == ["a", "b"]
    assert rows[0]["b"] == ""
    assert rows[1]["b"] == "3"


def test_write_keep_and_drop_raises(simple_csv: Path, tmp_path: Path) -> None:
    out = tmp_path / "out.csv"
    runner = ThaCSV()
    runner.read(None, simple_csv, ["name"])
    with pytest.raises(ValueError, match="Cannot specify both"):
        runner.write(None, out, keep=["name"], drop=["email"])


def test_write_column_order(simple_csv: Path, tmp_path: Path) -> None:
    out = tmp_path / "out.csv"
    runner = ThaCSV()
    runner.read(None, simple_csv, ["name"])
    runner.write(None, out, column_order=["email", "name"])
    rows = list(csv.DictReader(out.open()))
    keys = list(rows[0].keys())
    assert keys[0] == "email"
    assert keys[1] == "name"
    assert "id" in keys


def test_write_column_order_unlisted_follow(simple_csv: Path, tmp_path: Path) -> None:
    out = tmp_path / "out.csv"
    runner = ThaCSV()
    runner.read(None, simple_csv, ["name"])
    runner.write(None, out, column_order=["row number"])
    keys = next(iter(csv.DictReader(out.open()))).keys()
    assert next(iter(keys)) == "row number"


def test_sort_tie_preserves_order(tmp_path: Path) -> None:
    csv_path = tmp_path / "tied.csv"
    csv_path.write_text("id,name\n1,Alice\n2,Alice\n3,Alice\n")
    out = tmp_path / "out.csv"
    runner = ThaCSV()
    runner.read(None, csv_path, ["name"])
    runner.write(None, out, sort_by="name")
    rows = list(csv.DictReader(out.open()))
    assert [r["id"] for r in rows] == ["1", "2", "3"]


def test_sort_numeric_aware(tmp_path: Path) -> None:
    csv_path = tmp_path / "mixed.csv"
    csv_path.write_text("id,val\n1,10\n2,5\n3,abc\n")
    out = tmp_path / "out.csv"
    runner = ThaCSV()
    runner.read(None, csv_path, ["val"])
    runner.write(None, out, sort_by="val")
    rows = list(csv.DictReader(out.open()))
    vals = [r["val"] for r in rows]
    assert vals == ["5", "10", "abc"]


def test_desc_used_verbatim_as_read_label(
    simple_csv: Path, capsys: pytest.CaptureFixture[str]
) -> None:
    runner = ThaCSV()
    runner.read("[1/7]: Loading things", simple_csv, ["name"])
    err = capsys.readouterr().err
    assert runner.rows  # read completed
    assert "[1/7]: Loading things" in err
    assert "Reading" not in err


def test_desc_none_uses_default_read_label(
    simple_csv: Path, capsys: pytest.CaptureFixture[str]
) -> None:
    ThaCSV().read(None, simple_csv, ["name"])
    assert f"Reading {simple_csv.stem} CSV" in capsys.readouterr().err


def test_enrich_false_omits_enriched_columns(simple_csv: Path) -> None:
    runner = ThaCSV()
    runner.read(None, simple_csv, ["name"], enrich=False)
    for row in runner.rows:
        assert "row number" not in row
        assert "row status" not in row
        assert "message" not in row


def test_enrich_false_preserves_original_columns(simple_csv: Path) -> None:
    runner = ThaCSV()
    runner.read(None, simple_csv, ["name", "email"], enrich=False)
    assert runner.rows[0]["name"] == "Alice"
    assert runner.rows[0]["email"] == "alice@example.com"


def test_enrich_false_validator_error_still_raises(simple_csv: Path) -> None:
    runner = ThaCSV()
    with pytest.raises(ValueError, match="Bob is not allowed"):
        runner.read(None, simple_csv, ["name"], fail_on_bob, enrich=False)


# --- chunk_size ---


def test_chunk_size_returns_list(simple_csv: Path, tmp_path: Path) -> None:
    runner = ThaCSV()
    runner.read(None, simple_csv, ["name"])
    result = runner.write(None, tmp_path / "out.csv", chunk_size=2)
    assert isinstance(result, list)


def test_chunk_size_correct_file_count(simple_csv: Path, tmp_path: Path) -> None:
    runner = ThaCSV()
    runner.read(None, simple_csv, ["name"])
    paths = runner.write(None, tmp_path / "out.csv", chunk_size=2)
    assert isinstance(paths, list)
    assert len(paths) == 2  # 3 rows → chunks of 2, 1


def test_chunk_size_files_exist(simple_csv: Path, tmp_path: Path) -> None:
    runner = ThaCSV()
    runner.read(None, simple_csv, ["name"])
    paths = runner.write(None, tmp_path / "out.csv", chunk_size=2)
    assert isinstance(paths, list)
    assert all(p.exists() for p in paths)


def test_chunk_size_naming(simple_csv: Path, tmp_path: Path) -> None:
    runner = ThaCSV()
    runner.read(None, simple_csv, ["name"])
    paths = runner.write(None, tmp_path / "out.csv", chunk_size=2)
    assert isinstance(paths, list)
    assert paths[0].name == "out_001.csv"
    assert paths[1].name == "out_002.csv"


def test_chunk_size_total_rows(simple_csv: Path, tmp_path: Path) -> None:
    runner = ThaCSV()
    runner.read(None, simple_csv, ["name"])
    paths = runner.write(None, tmp_path / "out.csv", chunk_size=2)
    assert isinstance(paths, list)
    total = sum(len(list(csv.DictReader(p.open()))) for p in paths)
    assert total == 3


def test_chunk_size_larger_than_rows(simple_csv: Path, tmp_path: Path) -> None:
    runner = ThaCSV()
    runner.read(None, simple_csv, ["name"])
    paths = runner.write(None, tmp_path / "out.csv", chunk_size=100)
    assert isinstance(paths, list)
    assert len(paths) == 1


def test_chunk_size_zero_raises(simple_csv: Path, tmp_path: Path) -> None:
    runner = ThaCSV()
    runner.read(None, simple_csv, ["name"])
    with pytest.raises(ValueError, match="chunk_size"):
        runner.write(None, tmp_path / "out.csv", chunk_size=0)


def test_no_chunk_size_returns_path(simple_csv: Path, tmp_path: Path) -> None:
    runner = ThaCSV()
    runner.read(None, simple_csv, ["name"])
    result = runner.write(None, tmp_path / "out.csv")
    assert isinstance(result, Path)


# --- delimiter / TSV ---


def test_tsv_read(simple_tsv: Path) -> None:
    runner = ThaCSV(delimiter="\t")
    runner.read(None, simple_tsv, ["id", "name", "email"])
    assert len(runner.rows) == 3
    assert runner.rows[0]["name"] == "Alice"


def test_tsv_write(simple_tsv: Path, tmp_path: Path) -> None:
    out = tmp_path / "out.tsv"
    runner = ThaCSV(delimiter="\t")
    runner.read(None, simple_tsv, ["name"])
    runner.write(None, out)
    rows = list(csv.DictReader(out.open(), delimiter="\t"))
    assert len(rows) == 3
    assert rows[1]["name"] == "Bob"


def test_tsv_roundtrip_preserves_values(simple_tsv: Path, tmp_path: Path) -> None:
    out = tmp_path / "out.tsv"
    runner = ThaCSV(delimiter="\t")
    runner.read(None, simple_tsv, ["name", "email"], enrich=False)
    runner.write(None, out)
    rows = list(csv.DictReader(out.open(), delimiter="\t"))
    assert rows[0] == {"name": "Alice", "email": "alice@example.com", "id": "1"}


def test_default_delimiter_is_comma(simple_csv: Path) -> None:
    runner = ThaCSV()
    runner.read(None, simple_csv, ["name"])
    assert runner.rows[0]["name"] == "Alice"


# --- encoding ---


def test_encoding_read_cp1252(tmp_path: Path) -> None:
    csv_path = tmp_path / "encoded.csv"
    csv_path.write_bytes("name,city\nJos\xe9,S\xe3o Paulo\n".encode("cp1252"))
    runner = ThaCSV(encoding="cp1252")
    runner.read(None, csv_path, ["name"])
    assert runner.rows[0]["name"] == "José"
    assert runner.rows[0]["city"] == "São Paulo"


def test_encoding_write_cp1252(tmp_path: Path) -> None:
    csv_path = tmp_path / "encoded.csv"
    csv_path.write_bytes("name,city\nJos\xe9,S\xe3o Paulo\n".encode("cp1252"))
    out = tmp_path / "out.csv"
    runner = ThaCSV(encoding="cp1252")
    runner.read(None, csv_path, ["name"])
    runner.write(None, out)
    content = out.read_bytes().decode("cp1252")
    assert "José" in content
    assert "São Paulo" in content


def test_encoding_roundtrip_cp1252(tmp_path: Path) -> None:
    csv_path = tmp_path / "encoded.csv"
    csv_path.write_bytes("name,city\nJos\xe9,S\xe3o Paulo\n".encode("cp1252"))
    out = tmp_path / "out.csv"
    runner = ThaCSV(encoding="cp1252")
    runner.read(None, csv_path, ["name", "city"], enrich=False)
    runner.write(None, out)
    rows = list(csv.DictReader(out.open(encoding="cp1252")))
    assert rows[0]["name"] == "José"
    assert rows[0]["city"] == "São Paulo"


# --- excel ---


def test_write_xlsx(simple_csv: Path, tmp_path: Path) -> None:
    import openpyxl

    out = tmp_path / "out.xlsx"
    runner = ThaCSV()
    runner.read(None, simple_csv, ["id", "name", "email"], enrich=False)
    result = runner.write(None, out)

    assert result == out
    assert out.exists()

    wb = openpyxl.load_workbook(out)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    assert rows[0] == ("id", "name", "email")
    assert rows[1] == ("1", "Alice", "alice@example.com")


def test_read_xlsx(tmp_path: Path) -> None:
    import openpyxl

    xlsx_path = tmp_path / "input.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["name", "age"])
    ws.append(["Alice", 30])
    ws.append(["Bob", 25])
    wb.save(xlsx_path)

    runner = ThaCSV()
    rows = runner.read(None, xlsx_path, ["name", "age"], enrich=False)

    assert rows == [{"name": "Alice", "age": 30}, {"name": "Bob", "age": 25}]


def test_read_xlsx_missing_headers_raises(tmp_path: Path) -> None:
    import openpyxl

    xlsx_path = tmp_path / "input.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["name"])
    ws.append(["Alice"])
    wb.save(xlsx_path)

    runner = ThaCSV()
    with pytest.raises(CsvError, match="Missing required headers"):
        runner.read(None, xlsx_path, ["name", "email"])


def test_read_xlsx_empty_raises(tmp_path: Path) -> None:
    import openpyxl

    xlsx_path = tmp_path / "empty.xlsx"
    wb = openpyxl.Workbook()
    wb.save(xlsx_path)
    # default new workbook has one truly empty sheet with no rows at all
    ws = wb.active
    assert ws.max_row == 1

    runner = ThaCSV()
    with pytest.raises(CsvError, match="appears to be empty"):
        runner.read(None, xlsx_path, ["name"])


def test_xlsx_roundtrip(simple_csv: Path, tmp_path: Path) -> None:
    out = tmp_path / "out.xlsx"
    runner = ThaCSV()
    runner.read(None, simple_csv, ["id", "name", "email"], enrich=False)
    runner.write(None, out)

    reader = ThaCSV()
    rows = reader.read(None, out, ["id", "name", "email"], enrich=False)
    assert rows[0]["name"] == "Alice"


def test_xlsx_chunked_write(simple_csv: Path, tmp_path: Path) -> None:
    out = tmp_path / "out.xlsx"
    runner = ThaCSV()
    runner.read(None, simple_csv, ["id", "name", "email"], enrich=False)
    paths = runner.write(None, out, chunk_size=1)

    assert isinstance(paths, list)
    assert len(paths) == 3
    assert all(p.suffix == ".xlsx" for p in paths)
    assert all(p.exists() for p in paths)


def test_read_xlsx_sheet_by_name(tmp_path: Path) -> None:
    import openpyxl

    xlsx_path = tmp_path / "input.xlsx"
    wb = openpyxl.Workbook()
    wb.active.title = "First"
    wb.active.append(["name"])
    wb.active.append(["Alice"])
    second = wb.create_sheet("Second")
    second.append(["name"])
    second.append(["Bob"])
    wb.save(xlsx_path)

    runner = ThaCSV()
    rows = runner.read(None, xlsx_path, ["name"], enrich=False, sheet="Second")

    assert rows == [{"name": "Bob"}]


def test_read_xlsx_sheet_by_index(tmp_path: Path) -> None:
    import openpyxl

    xlsx_path = tmp_path / "input.xlsx"
    wb = openpyxl.Workbook()
    wb.active.append(["name"])
    wb.active.append(["Alice"])
    second = wb.create_sheet("Second")
    second.append(["name"])
    second.append(["Bob"])
    wb.save(xlsx_path)

    runner = ThaCSV()
    rows = runner.read(None, xlsx_path, ["name"], enrich=False, sheet=1)

    assert rows == [{"name": "Bob"}]


def test_read_xlsx_sheet_name_not_found_raises(tmp_path: Path) -> None:
    import openpyxl

    xlsx_path = tmp_path / "input.xlsx"
    wb = openpyxl.Workbook()
    wb.active.append(["name"])
    wb.save(xlsx_path)

    runner = ThaCSV()
    with pytest.raises(CsvError, match="Sheet 'Nope' not found"):
        runner.read(None, xlsx_path, ["name"], sheet="Nope")


def test_read_xlsx_sheet_index_out_of_range_raises(tmp_path: Path) -> None:
    import openpyxl

    xlsx_path = tmp_path / "input.xlsx"
    wb = openpyxl.Workbook()
    wb.active.append(["name"])
    wb.save(xlsx_path)

    runner = ThaCSV()
    with pytest.raises(CsvError, match="Sheet index 5 out of range"):
        runner.read(None, xlsx_path, ["name"], sheet=5)


def test_read_csv_with_sheet_raises(simple_csv: Path) -> None:
    runner = ThaCSV()
    with pytest.raises(ValueError, match="sheet= is only valid"):
        runner.read(None, simple_csv, ["name"], sheet="Sheet1")


def test_write_xlsx_names_sheet(simple_csv: Path, tmp_path: Path) -> None:
    import openpyxl

    out = tmp_path / "out.xlsx"
    runner = ThaCSV()
    runner.read(None, simple_csv, ["id", "name", "email"], enrich=False)
    runner.write(None, out, sheet="Contacts")

    wb = openpyxl.load_workbook(out)
    assert wb.sheetnames == ["Contacts"]


def test_write_xlsx_names_sheet_chunked(simple_csv: Path, tmp_path: Path) -> None:
    import openpyxl

    out = tmp_path / "out.xlsx"
    runner = ThaCSV()
    runner.read(None, simple_csv, ["id", "name", "email"], enrich=False)
    paths = runner.write(None, out, chunk_size=1, sheet="Contacts")

    assert isinstance(paths, list)
    for p in paths:
        wb = openpyxl.load_workbook(p)
        assert wb.sheetnames == ["Contacts"]


def test_write_csv_with_sheet_raises(simple_csv: Path, tmp_path: Path) -> None:
    out = tmp_path / "out.csv"
    runner = ThaCSV()
    runner.read(None, simple_csv, ["id", "name", "email"], enrich=False)
    with pytest.raises(ValueError, match="sheet= is only valid"):
        runner.write(None, out, sheet="Contacts")


# --- jsonl ---


def test_write_jsonl(simple_csv: Path, tmp_path: Path) -> None:
    import json

    out = tmp_path / "out.jsonl"
    runner = ThaCSV()
    runner.read(None, simple_csv, ["id", "name", "email"], enrich=False)
    result = runner.write(None, out)

    assert result == out
    lines = out.read_text().strip().splitlines()
    assert len(lines) == 3
    assert json.loads(lines[0]) == {"id": "1", "name": "Alice", "email": "alice@example.com"}


def test_read_jsonl(tmp_path: Path) -> None:
    import json

    jsonl_path = tmp_path / "input.jsonl"
    jsonl_path.write_text(
        json.dumps({"name": "Alice", "age": 30})
        + "\n"
        + json.dumps({"name": "Bob", "age": 25})
        + "\n"
    )

    runner = ThaCSV()
    rows = runner.read(None, jsonl_path, ["name", "age"], enrich=False)

    assert rows == [{"name": "Alice", "age": 30}, {"name": "Bob", "age": 25}]


def test_read_jsonl_skips_blank_lines(tmp_path: Path) -> None:
    import json

    jsonl_path = tmp_path / "input.jsonl"
    jsonl_path.write_text(json.dumps({"name": "Alice"}) + "\n\n\n" + json.dumps({"name": "Bob"}))

    runner = ThaCSV()
    rows = runner.read(None, jsonl_path, ["name"], enrich=False)

    assert rows == [{"name": "Alice"}, {"name": "Bob"}]


def test_read_jsonl_missing_headers_raises(tmp_path: Path) -> None:
    import json

    jsonl_path = tmp_path / "input.jsonl"
    jsonl_path.write_text(json.dumps({"name": "Alice"}) + "\n")

    runner = ThaCSV()
    with pytest.raises(CsvError, match="Missing required headers"):
        runner.read(None, jsonl_path, ["name", "email"])


def test_read_jsonl_empty_raises(tmp_path: Path) -> None:
    jsonl_path = tmp_path / "empty.jsonl"
    jsonl_path.write_text("")

    runner = ThaCSV()
    with pytest.raises(CsvError, match="appears to be empty"):
        runner.read(None, jsonl_path, ["name"])


def test_jsonl_roundtrip(simple_csv: Path, tmp_path: Path) -> None:
    out = tmp_path / "out.jsonl"
    runner = ThaCSV()
    runner.read(None, simple_csv, ["id", "name", "email"], enrich=False)
    runner.write(None, out)

    reader = ThaCSV()
    rows = reader.read(None, out, ["id", "name", "email"], enrich=False)
    assert rows[0]["name"] == "Alice"


def test_jsonl_chunked_write(simple_csv: Path, tmp_path: Path) -> None:
    out = tmp_path / "out.jsonl"
    runner = ThaCSV()
    runner.read(None, simple_csv, ["id", "name", "email"], enrich=False)
    paths = runner.write(None, out, chunk_size=1)

    assert isinstance(paths, list)
    assert len(paths) == 3
    assert all(p.suffix == ".jsonl" for p in paths)
    assert all(p.exists() for p in paths)


def test_jsonl_respects_column_filtering(simple_csv: Path, tmp_path: Path) -> None:
    import json

    out = tmp_path / "out.jsonl"
    runner = ThaCSV()
    runner.read(None, simple_csv, ["id", "name", "email"], enrich=False)
    runner.write(None, out, keep=["name"])

    lines = out.read_text().strip().splitlines()
    assert json.loads(lines[0]) == {"name": "Alice"}


def test_read_jsonl_with_sheet_raises(tmp_path: Path) -> None:
    import json

    jsonl_path = tmp_path / "input.jsonl"
    jsonl_path.write_text(json.dumps({"name": "Alice"}) + "\n")

    runner = ThaCSV()
    with pytest.raises(ValueError, match="sheet= is only valid"):
        runner.read(None, jsonl_path, ["name"], sheet="Sheet1")


def test_write_jsonl_with_sheet_raises(simple_csv: Path, tmp_path: Path) -> None:
    out = tmp_path / "out.jsonl"
    runner = ThaCSV()
    runner.read(None, simple_csv, ["id", "name", "email"], enrich=False)
    with pytest.raises(ValueError, match="sheet= is only valid"):
        runner.write(None, out, sheet="Sheet1")


# --- show_progress ---


def test_show_progress_false_suppresses_read_output(
    simple_csv: Path, capsys: pytest.CaptureFixture[str]
) -> None:
    runner = ThaCSV(show_progress=False)
    runner.read(None, simple_csv, ["name"])
    assert capsys.readouterr().err == ""


def test_show_progress_default_shows_read_output(
    simple_csv: Path, capsys: pytest.CaptureFixture[str]
) -> None:
    runner = ThaCSV()
    runner.read(None, simple_csv, ["name"])
    assert capsys.readouterr().err != ""


def test_show_progress_false_suppresses_write_output(
    simple_csv: Path, tmp_path: Path, capsys: pytest.CaptureFixture[str]
) -> None:
    out = tmp_path / "out.csv"
    runner = ThaCSV(show_progress=False)
    runner.read(None, simple_csv, ["name"], enrich=False)
    capsys.readouterr()  # discard read()'s output
    runner.write(None, out)
    assert capsys.readouterr().err == ""


def test_show_progress_false_suppresses_xlsx_write_output(
    simple_csv: Path, tmp_path: Path, capsys: pytest.CaptureFixture[str]
) -> None:
    out = tmp_path / "out.xlsx"
    runner = ThaCSV(show_progress=False)
    runner.read(None, simple_csv, ["name"], enrich=False)
    capsys.readouterr()
    runner.write(None, out)
    assert capsys.readouterr().err == ""


def test_show_progress_false_suppresses_jsonl_write_output(
    simple_csv: Path, tmp_path: Path, capsys: pytest.CaptureFixture[str]
) -> None:
    out = tmp_path / "out.jsonl"
    runner = ThaCSV(show_progress=False)
    runner.read(None, simple_csv, ["name"], enrich=False)
    capsys.readouterr()
    runner.write(None, out)
    assert capsys.readouterr().err == ""


# --- progress label / bar format / Done spacing ---


def _messages(runner: ThaCSV) -> list[str]:
    messages: list[str] = []
    runner.status_cb = messages.append
    return messages


def test_write_desc_used_verbatim_as_label(
    simple_csv: Path, tmp_path: Path, capsys: pytest.CaptureFixture[str]
) -> None:
    runner = ThaCSV()
    runner.read(None, simple_csv, ["name"], enrich=False)
    capsys.readouterr()
    runner.write("[6/7]: Saving results", tmp_path / "out.csv")
    err = capsys.readouterr().err
    assert "[6/7]: Saving results" in err
    assert "Writing" not in err


def test_write_desc_none_uses_default_label(
    simple_csv: Path, tmp_path: Path, capsys: pytest.CaptureFixture[str]
) -> None:
    runner = ThaCSV()
    runner.read(None, simple_csv, ["name"], enrich=False)
    capsys.readouterr()
    runner.write(None, tmp_path / "out.csv")
    assert "Writing out CSV" in capsys.readouterr().err


def test_chunked_write_label_appends_chunk_index_to_desc(
    simple_csv: Path, tmp_path: Path, capsys: pytest.CaptureFixture[str]
) -> None:
    runner = ThaCSV()
    runner.read(None, simple_csv, ["name"], enrich=False)
    capsys.readouterr()
    runner.write("Saving", tmp_path / "out.csv", chunk_size=1)
    err = capsys.readouterr().err
    assert "Saving (1/" in err
    assert "Writing" not in err


def test_chunked_write_desc_none_uses_default_label(
    simple_csv: Path, tmp_path: Path, capsys: pytest.CaptureFixture[str]
) -> None:
    runner = ThaCSV()
    runner.read(None, simple_csv, ["name"], enrich=False)
    capsys.readouterr()
    runner.write(None, tmp_path / "out.csv", chunk_size=1)
    assert "Writing out CSV (1/" in capsys.readouterr().err


def test_bar_omits_elapsed_and_rate(
    simple_csv: Path, tmp_path: Path, capsys: pytest.CaptureFixture[str]
) -> None:
    runner = ThaCSV()
    runner.read("Reading", simple_csv, ["name"])
    runner.write("Writing", tmp_path / "out.csv")
    err = capsys.readouterr().err
    assert "it/s" not in err
    assert "[00:" not in err
    assert "100%" in err


def test_done_line_has_leading_blank_line_when_progress_shown(
    simple_csv: Path, tmp_path: Path
) -> None:
    runner = ThaCSV()
    messages = _messages(runner)
    runner.read(None, simple_csv, ["name"], enrich=False)
    runner.write(None, tmp_path / "out.csv")
    assert messages[-1].startswith("\n✅ Done! CSV was written to: ")


def test_done_line_has_leading_blank_line_when_chunked(simple_csv: Path, tmp_path: Path) -> None:
    runner = ThaCSV()
    messages = _messages(runner)
    runner.read(None, simple_csv, ["name"], enrich=False)
    runner.write(None, tmp_path / "out.csv", chunk_size=1)
    assert messages[-1].startswith("\n✅ Done! CSV was written to: ")


def test_done_line_has_no_blank_line_when_progress_hidden(simple_csv: Path, tmp_path: Path) -> None:
    runner = ThaCSV(show_progress=False)
    messages = _messages(runner)
    runner.read(None, simple_csv, ["name"], enrich=False)
    runner.write(None, tmp_path / "out.csv")
    assert messages[-1].startswith("✅ Done! CSV was written to: ")


def test_done_line_has_no_blank_line_when_no_rows(tmp_path: Path) -> None:
    runner = ThaCSV()
    messages = _messages(runner)
    runner.write(None, tmp_path / "out.csv", rows=[])
    assert messages[-1].startswith("✅ Done! CSV was written to: ")
